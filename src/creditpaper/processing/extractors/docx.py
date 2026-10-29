"""DOCX text extraction. Paragraphs + tables. Treated as one logical page."""

from __future__ import annotations

import io

from docx import Document

from creditpaper.processing.models import ExtractionResult


def extract_docx(raw: bytes, *, original_filename: str) -> ExtractionResult:
    doc = Document(io.BytesIO(raw))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    full_text = "\n\n".join(parts)
    return ExtractionResult(
        full_text=full_text,
        page_texts=[{"page": 1, "text": full_text, "ocr_used": False}],
        page_count=1,
    )
