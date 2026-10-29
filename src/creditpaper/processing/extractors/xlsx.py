"""XLSX extraction as flat row dump.

Each sheet becomes one logical "page" with `sheet_name` metadata.
Formula cells return cached values (`data_only=True`).
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from creditpaper.processing.models import ExtractionResult


def extract_xlsx(raw: bytes, *, original_filename: str) -> ExtractionResult:
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    parts: list[str] = []
    page_texts: list[dict] = []
    for s_idx, sheet in enumerate(wb.worksheets, start=1):
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append(" | ".join(cells))
        sheet_text = f"# Sheet: {sheet.title}\n" + "\n".join(rows)
        parts.append(sheet_text)
        page_texts.append(
            {
                "page": s_idx,
                "text": sheet_text,
                "ocr_used": False,
                "sheet_name": sheet.title,
            }
        )
    sheet_count = len(wb.worksheets)
    wb.close()
    full_text = "\n\n".join(parts)
    return ExtractionResult(
        full_text=full_text,
        page_texts=page_texts,
        page_count=None,
        sheet_count=sheet_count,
    )
