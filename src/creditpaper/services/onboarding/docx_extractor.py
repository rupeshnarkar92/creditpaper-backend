"""DOCX text / section / style extractors (python-docx, already a dependency).

Public surface:
  - extract_text(content: bytes) -> str
  - extract_sections(content: bytes) -> list[dict]  (same shape as the PDF one)
  - extract_style(content: bytes) -> dict           (same shape as the PDF one)

Nothing here raises — corrupt files yield "", [] or Nones so the caller can
still persist the document row with a FILE_CORRUPTED-style outcome upstream.
"""

from __future__ import annotations

import io
from typing import Any


def _open(content: bytes):
    try:
        from docx import Document

        return Document(io.BytesIO(content))
    except Exception:
        return None


def extract_text(content: bytes) -> str:
    doc = _open(content)
    if doc is None:
        return ""
    try:
        parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        parts.append(t)
        return "\n".join(parts)
    except Exception:
        return ""


def extract_sections(content: bytes) -> list[dict[str, Any]]:
    """Headings become sections: prefer Heading 1, fall back to Heading 2.

    Each section carries a `preview` (body text under the heading, capped)
    matching the PDF extractor's shape, and `subsections` from the next heading
    level down (Heading 2 under Heading 1, Heading 3 under Heading 2), numbered
    within the section ("1.1", "1.2"). `subsections` is omitted when there are none.
    """
    doc = _open(content)
    if doc is None:
        return []
    try:
        for level, sub_level in (("Heading 1", "Heading 2"), ("Heading 2", "Heading 3")):
            collected: list[dict[str, Any]] = []
            current: dict[str, Any] | None = None
            for p in doc.paragraphs:
                text = (p.text or "").strip()
                if not text:
                    continue
                if p.style.name == level:
                    if current is not None:
                        collected.append(current)
                    current = {"name": text, "body": [], "subs": []}
                elif current is not None:
                    if p.style.name == sub_level:
                        current["subs"].append(text)
                    current["body"].append(text)
            if current is not None:
                collected.append(current)
            if collected:
                out: list[dict[str, Any]] = []
                seen: set[str] = set()
                for c in collected:
                    if c["name"] in seen:
                        continue
                    seen.add(c["name"])
                    preview = " ".join(c["body"])[:600].rstrip()
                    section: dict[str, Any] = {
                        "number": f"{len(out) + 1:02d}",
                        "name": c["name"],
                        "preview": preview,
                    }
                    subsections = _number_subsections(len(out) + 1, c["subs"])
                    if subsections:
                        section["subsections"] = subsections
                    out.append(section)
                    if len(out) >= 10:
                        break
                if out:
                    return out
        return []
    except Exception:
        return []


def _number_subsections(section_number: int, names: list[str]) -> list[dict[str, str]]:
    """Dedupe (case- and space-insensitive) and number names within their section, capped at 30."""
    out: list[dict[str, str]] = []
    seen: set[str] = set()
    for name in names:
        key = " ".join(name.lower().split())
        if not key or key in seen:
            continue
        seen.add(key)
        out.append({"number": f"{section_number}.{len(out) + 1}", "name": name})
        if len(out) >= 30:
            break
    return out


def extract_style(content: bytes) -> dict[str, Any]:
    """Heading/body font names from the document styles. Colors from raw
    python-docx runs are theme-dependent and unreliable — left None."""
    out: dict[str, Any] = {"font_heading": None, "font_body": None, "primary_color": None}
    doc = _open(content)
    if doc is None:
        return out
    try:
        heading = None
        body = None
        try:
            heading = doc.styles["Heading 1"].font.name
        except KeyError:
            heading = None
        try:
            body = doc.styles["Normal"].font.name
        except KeyError:
            body = None
        out["font_heading"] = heading
        out["font_body"] = body or heading
    except Exception:
        pass
    return out
