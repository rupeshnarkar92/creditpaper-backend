"""Paper specification — the paper-level format contract.

Formatting is a PAPER-level concern, not a section-level one: numbering,
typography, heading hierarchy, fonts, colours, margins, and page numbers are
defined once here and applied uniformly by the renderer — the formatting analogue
of the agent constitution. Visual identity per client arrives as a resolved paper
theme (fonts, colours, logo, standard typography scale; D10) through
`PaperSpec.from_theme`; `DEFAULT_SPEC` is what a render without a theme uses.

Numbering is applied manually by the renderer (deterministic 1 / 1.1 / 1.1.1)
rather than via Word's fragile numbering XML.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, fields
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

# Image types python-docx can embed inline. SVG and WebP logos are served to the web
# paper but left off the .docx (the cover then carries the firm name only).
DOCX_IMAGE_TYPES = frozenset({"image/png", "image/jpeg", "image/gif"})


@dataclass(frozen=True)
class PaperSpec:
    body_font: str = "Calibri"
    heading_font: str = "Calibri Light"
    body_size: float = 10.0
    h1_size: float = 16.0
    h2_size: float = 12.5
    h3_size: float = 11.0
    heading_rgb: tuple[int, int, int] = (0x14, 0x3C, 0x5A)  # deep navy
    body_rgb: tuple[int, int, int] = (0x22, 0x22, 0x22)
    draft_rgb: tuple[int, int, int] = (0xB0, 0x10, 0x10)  # red — draft / placeholders
    muted_rgb: tuple[int, int, int] = (0x70, 0x70, 0x70)
    ref_rgb: tuple[int, int, int] = (0x3A, 0x6E, 0xA5)  # provenance refs (superscript)
    total_shade_hex: str = "DCE3EC"  # Total column / totals row fill
    margin_cm: float = 2.0
    page_numbers: bool = True
    # Paper theme (spec 2026-09-15 §4). The defaults leave the renderer's existing
    # output unchanged; `from_theme` fills them from a resolved theme.
    cover_title_size: float | None = None
    table_size: float | None = None
    caption_size: float = 8.0
    line_spacing: float | None = None
    page_size: str | None = None  # "A4"; None keeps python-docx's default page
    running_header: bool = True
    footer_note: str | None = None
    accent_rgb: tuple[int, int, int] | None = None
    firm_name: str | None = None
    logo_bytes: bytes | None = None
    theme_version: str | None = None

    @classmethod
    def from_theme(cls, theme: dict[str, Any], logo_bytes: bytes | None = None) -> PaperSpec:
        """A PaperSpec from a resolved paper theme (the `GET /paper/theme` payload).

        Fonts use each family's docx fallback name (a face Word reliably has); sizes
        come from the typography scale; heading, text, muted and accent colours from
        the theme. Missing or malformed values keep the dataclass default, so a
        partial theme still renders. Output stays structurally plain (D2).
        """
        base = cls()
        fonts = theme.get("fonts") or {}
        colours = theme.get("colours") or {}
        typo = theme.get("typography") or {}
        page = theme.get("page") or {}

        def font(role: str, default: str) -> str:
            value = (fonts.get(role) or {}).get("docx")
            return value if isinstance(value, str) and value.strip() else default

        def size(key: str, default: float | None) -> float | None:
            value = typo.get(key)
            return float(value) if isinstance(value, int | float) and not isinstance(value, bool) else default

        def rgb(key: str, default: tuple[int, int, int] | None) -> tuple[int, int, int] | None:
            return _parse_rgb(colours.get(key)) or default

        tint = colours.get("accent_tint")
        margin = page.get("margin_mm")
        values: dict[str, Any] = {
            "body_font": font("body", base.body_font),
            "heading_font": font("heading", base.heading_font),
            "body_size": size("body_pt", base.body_size),
            "h1_size": size("h1_pt", base.h1_size),
            "h2_size": size("h2_pt", base.h2_size),
            "h3_size": size("h3_pt", base.h3_size),
            "cover_title_size": size("cover_title_pt", None),
            "table_size": size("table_pt", None),
            "caption_size": size("caption_pt", base.caption_size),
            "line_spacing": size("line_height", None),
            "heading_rgb": rgb("heading", base.heading_rgb),
            "body_rgb": rgb("text", base.body_rgb),
            "muted_rgb": rgb("muted", base.muted_rgb),
            "ref_rgb": rgb("accent", base.ref_rgb),
            "accent_rgb": rgb("accent", None),
            "total_shade_hex": tint[1:].upper() if _parse_rgb(tint) else base.total_shade_hex,
            "margin_cm": float(margin) / 10
            if isinstance(margin, int | float) and margin > 0
            else base.margin_cm,
            "page_numbers": bool(page.get("page_numbers", base.page_numbers)),
            "running_header": bool(page.get("running_header", base.running_header)),
            "page_size": page.get("size") if page.get("size") in _PAGE_SIZES else None,
            "footer_note": str(page["footer_note"]) if page.get("footer_note") else None,
            "firm_name": str(theme["firm_name"]) if theme.get("firm_name") else None,
            "logo_bytes": logo_bytes,
            "theme_version": str(theme["version"]) if theme.get("version") else None,
        }
        known = {f.name for f in fields(cls)}
        return cls(**{k: v for k, v in values.items() if k in known})


DEFAULT_SPEC = PaperSpec()

_PAGE_SIZES = {"A4": (Mm(210), Mm(297)), "Letter": (Mm(215.9), Mm(279.4))}
_HEX = re.compile(r"^#([0-9a-fA-F]{6})$")


def _parse_rgb(value: object) -> tuple[int, int, int] | None:
    if not isinstance(value, str):
        return None
    m = _HEX.match(value.strip())
    if not m:
        return None
    h = m.group(1)
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _rgb(t: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*t)


def _page_number_footer(doc, spec: PaperSpec) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    if spec.footer_note:
        note = p.add_run(f"{spec.footer_note}    ")
        note.font.size = Pt(spec.caption_size)
        note.font.color.rgb = _rgb(spec.muted_rgb)
    run = p.add_run("Page ")
    run.font.size = Pt(spec.caption_size)
    run.font.color.rgb = _rgb(spec.muted_rgb)
    # PAGE field
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = p.add_run()
    r.font.size = Pt(spec.caption_size)
    r.font.color.rgb = _rgb(spec.muted_rgb)
    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_end)


def apply_paper_spec(doc, spec: PaperSpec = DEFAULT_SPEC) -> None:
    """Set document-wide styles, margins, and page numbers from the spec."""
    normal = doc.styles["Normal"]
    normal.font.name = spec.body_font
    normal.font.size = Pt(spec.body_size)
    normal.font.color.rgb = _rgb(spec.body_rgb)
    if spec.line_spacing:
        normal.paragraph_format.line_spacing = spec.line_spacing

    for level, size in ((1, spec.h1_size), (2, spec.h2_size), (3, spec.h3_size)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = spec.heading_font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _rgb(spec.heading_rgb)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(4)

    title = doc.styles["Title"]
    title.font.name = spec.heading_font
    title.font.color.rgb = _rgb(spec.heading_rgb)
    if spec.cover_title_size:
        title.font.size = Pt(spec.cover_title_size)

    for section in doc.sections:
        if spec.page_size in _PAGE_SIZES:
            section.page_width, section.page_height = _PAGE_SIZES[spec.page_size]
        section.top_margin = Cm(spec.margin_cm)
        section.bottom_margin = Cm(spec.margin_cm)
        section.left_margin = Cm(spec.margin_cm)
        section.right_margin = Cm(spec.margin_cm)

    if spec.page_numbers:
        _page_number_footer(doc, spec)
    elif spec.footer_note:
        footer = doc.sections[0].footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(spec.footer_note)
        run.font.size = Pt(spec.caption_size)
        run.font.color.rgb = _rgb(spec.muted_rgb)
