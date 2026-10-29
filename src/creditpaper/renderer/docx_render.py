"""Word (.docx) renderer — deliberately plain, shaped for Canva import (D2).

All formatting (numbering, typography, heading hierarchy, fonts, colours, page
numbers) is paper-level and comes from the PaperSpec (renderer/paper_spec.py),
applied uniformly here — the renderer never decides format per section. The
renderer is disposable (D1): it reads the model, applies the spec, and resolves
citations; it never computes deal facts.

Structure:
  - Headings are manually numbered (1, 1.1, 1.1.1) per the spec.
  - Sections 1 (Deal Snapshot) and 3 (Funding Snapshot) are native tables;
    Section 7 (Risks) is a native table parsed from the agent's structured prose.
  - Other sections are prose: agents mark subsection headings with "## " / "### ",
    which render as numbered H2 / H3 (not inline text).
  - Draft sections carry a visible [DRAFT] indicator; unpopulated content renders
    as a marked "[to be populated: …]" placeholder in the draft colour.
  - Provenance is grouped by source (Option A): per-document / per-formula refs.
"""

from __future__ import annotations

import calendar
import io
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from creditpaper.format import EM_DASH, money, percent_of
from creditpaper.markup import ANALYST_PREFIX, SCAFFOLD_PREFIX
from creditpaper.model.citations import resolve_path
from creditpaper.model.core import (
    ExtractionMethod,
    Fact,
    SectionStatus,
    machine_recommendations,
)
from creditpaper.model.deal import DealModel
from creditpaper.renderer.paper_spec import DEFAULT_SPEC, PaperSpec, apply_paper_spec

# Fallback titles only. The renderer previously imported one tenant's section list
# from `profiles` and used it unconditionally, so a second lender's paper would
# have rendered under 3Capital's headings — and it pulled the whole
# renderer -> profiles -> agents chain along with it. Titles now arrive from the
# caller's house style; these are what a paper falls back to when none is given.
_DEFAULT_SECTION_TITLES: dict[int, str] = {
    1: "Deal Snapshot",
    2: "Deal Summary",
    3: "Funding Snapshot",
    4: "Zoning & Location",
    5: "Borrower Analysis",
    6: "Valuation Analysis",
    7: "Risks & Mitigations",
    8: "Other Considerations",
    9: "Conclusion",
}


def _section_titles(house_style: dict | None) -> dict[int, str]:
    """Section titles from the caller's house style, else the defaults.

    Accepts the profile's `model_dump()` shape, which is what every caller
    already has to hand — the alternative was importing the profile type here and
    reintroducing the dependency this change exists to remove.
    """
    sections = (house_style or {}).get("sections") or []
    titles = {
        s["section"]: str(s["title"]).title()
        for s in sections
        if isinstance(s, dict) and s.get("section") and s.get("title")
    }
    return {**_DEFAULT_SECTION_TITLES, **titles}


_FUNDING_ORDER = [
    "land",
    "land_transaction_costs",
    "construction",
    "contingency",
    "professional_fees",
    "stat_fees_and_contributions",
    "marketing",
    "land_holding_costs",
    "legals_settlement_costs",
    "broker_fee",
    "line_fee",
    "interest",
    "admin_fee",
    "management_fee",
    "establishment_fee",
    "spv_setup_cost",
]
_LABELS = {
    "land": "Land",
    "land_transaction_costs": "Land transaction costs",
    "construction": "Construction",
    "contingency": "Contingency",
    "professional_fees": "Professional fees",
    "stat_fees_and_contributions": "Stat. fees and contributions",
    "broker_fee": "Broker fee",
    "line_fee": "Line fee",
    "interest": "Interest",
    "establishment_fee": "Establishment fee",
    "spv_setup_cost": "SPV setup cost",
}


class _Provenance:
    """Markers grouped by source (Option A): same document/page share one ref;
    derived facts group by formula. Collapses per-figure refs to a handful."""

    def __init__(self) -> None:
        self._index: dict[tuple, int] = {}
        self.entries: list[tuple[int, object]] = []

    def ref_num(self, fact: Fact | None) -> int | None:
        if fact is None or fact.provenance is None or fact.value is None:
            return None
        p = fact.provenance
        key = ("derived", p.location) if p.method == ExtractionMethod.DERIVED else (p.document_id, p.page)
        if key not in self._index:
            self._index[key] = len(self._index) + 1
            self.entries.append((self._index[key], p))
        return self._index[key]

    def marker(self, fact: Fact | None) -> str:
        n = self.ref_num(fact)
        return f" [{n}]" if n else ""


class _WipRegister:
    """Outstanding items harvested during render — snapshot placeholders, scaffolded
    subsections, and explicit analyst actions — surfaced as one 'Items to be
    Confirmed' register so the paper's done-ness is unambiguous (Jason 2026-06-15)."""

    def __init__(self) -> None:
        self.items: list[tuple[str, str]] = []  # (section label, item)

    def add(self, section_label: str, item: str) -> None:
        item = " ".join(item.split())
        if item and (section_label, item) not in self.items:
            self.items.append((section_label, item))


def _money(fact: Fact | None) -> str:
    """Whole dollars via the shared formatter — see `creditpaper/format.py`.

    This used to format with `f"{value:,.0f}"`, i.e. Decimal's default
    ROUND_HALF_EVEN, while derive.py quantises ROUND_HALF_UP — so a column and its
    own total could round different ways and print a table that does not add up.
    """
    return money(fact.value if fact is not None else None, dash=EM_DASH)


def _fmt_date(d: date) -> str:
    return f"{d.day} {calendar.month_name[d.month]} {d.year}"


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _mark_header_row(row) -> None:
    """Mark a flat table's first row as a repeating semantic header for assistive tech."""
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _emit_value(
    paragraph,
    text: str,
    fact: Fact | None,
    prov: _Provenance,
    show: bool,
    spec: PaperSpec,
    bold: bool = False,
) -> None:
    """Write a value into a paragraph; newlines become line breaks; provenance ref
    as a small superscript run in the ref colour (declutters tables vs inline [n])."""
    for i, seg in enumerate(text.split("\n")):
        if i:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        paragraph.add_run(seg).bold = bold
    if show and fact is not None:
        n = prov.ref_num(fact)
        if n:
            ref = paragraph.add_run(str(n))
            ref.font.superscript = True
            ref.font.size = Pt(7)
            ref.font.color.rgb = RGBColor(*spec.ref_rgb)


def _cell_value(
    cell,
    text: str,
    fact: Fact | None,
    prov: _Provenance,
    show: bool,
    spec: PaperSpec,
    bold: bool = False,
    shade: str | None = None,
) -> None:
    _emit_value(cell.paragraphs[0], text, fact, prov, show, spec, bold)
    if shade:
        _shade(cell, shade)


def _placeholder(
    paragraph_or_cell,
    label: str,
    spec: PaperSpec,
    wip: _WipRegister | None = None,
    section_label: str | None = None,
) -> None:
    """Render a marked boilerplate placeholder in the draft colour, and register it
    as an outstanding item where a WIP register and section label are supplied."""
    run = (
        paragraph_or_cell.paragraphs[0].add_run(f"[to be populated: {label}]")
        if hasattr(paragraph_or_cell, "paragraphs")
        else paragraph_or_cell.add_run(f"[to be populated: {label}]")
    )
    run.italic = True
    run.font.color.rgb = RGBColor(*spec.draft_rgb)
    if wip is not None and section_label is not None:
        wip.add(section_label, f"Populate {label}")


_AUTHORSHIP_NOTES = {
    "ai_generated": "Drafted by the credit engine. Review and correct before approval.",
    "human_edited": "Drafted by the credit engine and edited by the analyst.",
    "human_authored": "Written by the analyst.",
    "deterministic": None,  # a table of anchored facts; nothing composed it
}


def _authorship_note(doc, authorship, spec: PaperSpec) -> None:
    """A one-line note under the heading saying who wrote the section."""
    note = _AUTHORSHIP_NOTES.get(getattr(authorship, "value", str(authorship)))
    if not note:
        return
    para = doc.add_paragraph()
    run = para.add_run(note)
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(*spec.ref_rgb)


def _heading(doc, text: str, level: int, spec: PaperSpec, draft: bool = False):
    h = doc.add_heading(text, level=level)
    if draft:
        tag = h.add_run("  [DRAFT]")
        tag.font.color.rgb = RGBColor(*spec.draft_rgb)
        tag.font.size = Pt(9)
        tag.bold = True
    return h


def _months_after(d: date, months: int) -> date:
    m = d.month - 1 + months
    return date(d.year + m // 12, m % 12 + 1, min(d.day, 28))


# ---------------------------------------------------------------------------
# Section renderers
# ---------------------------------------------------------------------------


# How each valuation basis is named in the paper. Keyed by the basis recorded on
# `metrics.lvr_basis`, so a figure and its label can never disagree about what the
# percentage is of.
_BASIS_LABELS = {
    "as_if_complete": "As If Complete (ex GST)",
    "as_if_complete_one_line": "As If Complete, in one line (ex GST)",
    "prsv": "Project-Related Site Value",
    "as_is": "As Is (ex GST)",
}


def _basis_label(basis_fact: Fact | None) -> str:
    """The paper's name for the basis an LVR was taken on."""
    key = str(basis_fact.value) if basis_fact is not None and basis_fact.value else ""
    return _BASIS_LABELS.get(key, key.replace("_", " ").title() or "basis to be confirmed")


def _pct_fact(fact: Fact | None, basis: str) -> str | None:
    """A derived percentage as 'NN.NN% of <basis>', via the shared formatter."""
    if fact is None or fact.value is None:
        return None
    return percent_of(fact.value, basis)


def _snapshot_rows(model: DealModel) -> list[tuple[str, str | None, Fact | None]]:
    """(label, value or None-for-placeholder, fact). Order (Jason 2026-06-15):
    context -> facility details -> sponsor/guarantors -> covenants -> exit.
    LVR rows cite model.metrics (one canonical figure, not a renderer recompute)."""
    e, v, mt = model.economics, model.valuation, model.metrics
    dtype = model.deal.type.value.value.replace("_", " ").title() if model.deal.type.value else None
    fac_type = f"First Mortgage {dtype} Facility" if dtype else None
    fund = model.deal.expected_funding_date.value
    maturity = None
    if fund and e.term_months.value:
        maturity = _fmt_date(_months_after(fund, int(e.term_months.value)))
    guarantors = "; ".join(str(f.value) for f in model.its.guarantors.values() if f.value) or None
    fees = [
        f"{fl.pct.value}% {key}"
        for key in ("establishment", "line")
        if (fl := e.fees.get(key)) and fl.pct.value is not None
    ]
    covs = "\n".join(f"• {c.summary.value}" for c in model.covenants.values() if c.summary.value) or None
    aic = v.values_by_basis.get("as_if_complete")
    # Security. This row used to be SYNTHESISED — "First mortgage over {address}"
    # assembled from the property address and the deal type, rendered with
    # fact=None so it carried no provenance marker at all. "First mortgage" is a
    # legal characterisation asserting RANKING, and nothing in the model
    # established it: `security.first_mortgage_available` exists in the schema but
    # no extractor populates it. So every paper stated an unsourced legal
    # conclusion about the security position, indistinguishable from an extracted
    # fact (D1/D4).
    #
    # Now: assert the ranking only where an anchored fact establishes it, and cite
    # that fact. Otherwise render a placeholder and register the gap. Guarantees
    # are not repeated here — they have their own row below.
    ranking = model.security.first_mortgage_available
    addr = v.property.address.value
    if ranking.value is True and addr:
        security, security_fact = f"First mortgage over {addr}", ranking
    elif ranking.value is False and addr:
        security, security_fact = f"Mortgage over {addr} (not first-ranking)", ranking
    else:
        security, security_fact = None, None
    return [
        # context
        ("Property", v.property.address.value, v.property.address),
        ("Facility type", fac_type, model.deal.type),
        ("Facility purpose", None, None),
        # facility details
        (
            "Total facility limit",
            _money(e.facility_amount) if e.facility_amount.value else None,
            e.facility_amount,
        ),
        (
            "Borrower interest rate",
            f"{e.interest_rate.value}% p.a." if e.interest_rate.value else None,
            e.interest_rate,
        ),
        ("Deal IRR", None, None),  # human-entered per paper; placeholder until confirmed
        ("Fees", "; ".join(fees) or None, None),
        ("Expected funding date", _fmt_date(fund) if fund else None, model.deal.expected_funding_date),
        ("Expected maturity", maturity, None),
        (
            "Valuation — As Is (ex GST)",
            _money(v.values_by_basis.get("as_is")) if v.values_by_basis.get("as_is") else None,
            v.values_by_basis.get("as_is"),
        ),
        ("Valuation — As If Complete (ex GST)", _money(aic) if aic else None, aic),
        # The ongoing LVR names the basis it was actually taken on. Hardcoding
        # "As If Complete" here would relabel a project-related-site-value LVR as
        # something it is not — the exact mislabelling the basis rule exists to
        # prevent (DECISIONS.md 2026-08-21).
        (
            f"Net facility LVR ({_basis_label(mt.lvr_basis)})",
            _pct_fact(mt.lvr_pct, _basis_label(mt.lvr_basis)),
            mt.lvr_pct,
        ),
        ("Net facility LVR (As Is, Day 1)", _pct_fact(mt.asis_lvr_pct, "As Is (ex GST)"), mt.asis_lvr_pct),
        ("Security", security, security_fact),
        # sponsor / guarantors
        ("Borrower", model.its.borrower_entity.value, model.its.borrower_entity),
        ("Sponsor / guarantors", guarantors, None),
        # covenants
        ("Financial covenants", covs, None),
        # exit
        ("Exit strategy", model.deal.exit_strategy.value, model.deal.exit_strategy),
    ]


def _render_snapshot(
    doc,
    model: DealModel,
    prov: _Provenance,
    show: bool,
    spec: PaperSpec,
    wip: _WipRegister | None = None,
    section_label: str | None = None,
) -> None:
    rows = _snapshot_rows(model)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text = "Item"
    t.rows[0].cells[1].text = "Value"
    _mark_header_row(t.rows[0])
    for cell in t.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for label, value, fact in rows:
        cells = t.add_row().cells
        cells[0].text = label
        if value is None:
            _placeholder(cells[1], label.lower(), spec, wip, section_label)
        else:
            _emit_value(cells[1].paragraphs[0], str(value), fact, prov, show, spec)


def _render_funding(doc, model: DealModel, prov: _Provenance, show: bool, spec: PaperSpec) -> None:
    e = model.economics
    ft = e.funding_table
    shade = spec.total_shade_hex
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    _mark_header_row(table.rows[0])
    for i, h in enumerate(("Item", "Equity", "Senior", "Total")):
        table.rows[0].cells[i].text = h
        table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        if i == 3:  # Total column header distinct
            _shade(table.rows[0].cells[i], shade)
    ordered = [k for k in _FUNDING_ORDER if k in ft] + [k for k in ft if k not in _FUNDING_ORDER]
    for key in ordered:
        line = ft[key]
        cells = table.add_row().cells
        cells[0].text = _LABELS.get(key, line.label or key.replace("_", " ").title())
        _cell_value(cells[1], _money(line.equity), line.equity, prov, show, spec)
        _cell_value(cells[2], _money(line.senior), line.senior, prov, show, spec)
        _cell_value(
            cells[3], _money(line.total), line.total, prov, show, spec, bold=True, shade=shade
        )  # Total column distinct
    totals = table.add_row().cells
    totals[0].text = "Total"
    totals[0].paragraphs[0].runs[0].bold = True
    _shade(totals[0], shade)
    _cell_value(totals[1], _money(e.equity_total), e.equity_total, prov, show, spec, bold=True, shade=shade)
    _cell_value(totals[2], _money(e.senior_total), e.senior_total, prov, show, spec, bold=True, shade=shade)
    _cell_value(totals[3], _money(e.total_cost), e.total_cost, prov, show, spec, bold=True, shade=shade)

    # LVR and LCR belong in the funding table (comment 30); cite model.metrics. The
    # ratio sits in the Total column (where the headline figures live), shaded.
    mt = model.metrics
    for label, fact, basis in (
        # Same rule as the snapshot: the label names the basis actually selected,
        # never a hardcoded one. The funding table and the snapshot must agree —
        # two rows describing one ratio under two different basis names is the
        # inconsistency the derive-once contract exists to prevent.
        (f"Net facility LVR ({_basis_label(mt.lvr_basis)})", mt.lvr_pct, _basis_label(mt.lvr_basis)),
        ("Loan to Cost Ratio (LCR)", mt.lcr_pct, "total cost"),
    ):
        if fact.value is None:
            continue
        cells = table.add_row().cells
        cells[0].text = label
        cells[0].paragraphs[0].runs[0].bold = True
        _cell_value(cells[3], percent_of(fact.value, basis), fact, prov, show, spec, bold=True, shade=shade)


def _parse_risks(text: str) -> list[dict]:
    """Parse '## name / Assessment: x / Mitigant: y' blocks. Scaffold ('>> ') and
    analyst ('[ANALYST] ') lines are captured separately, NOT folded into the
    mitigant — they predate the risk-table format and otherwise leak raw markup
    into the cell (review round 1)."""
    blocks: list[dict] = []
    cur: dict | None = None
    last_field: str | None = None  # continuation lines belong to the field they follow
    orphaned: list[str] = []  # content that arrived before any risk heading
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("## "):
            if cur:
                blocks.append(cur)
            cur = {
                "name": s[3:].strip(),
                "assessment": "",
                "mitigant": "",
                "scaffold": [],
                "analyst": [],
                "unparsed": [],
            }
            last_field = None
        elif cur is None:
            # Content before the first risk heading. Previously dropped outright,
            # so a section written as "## Key Risks" + "### <risk>" lost every
            # line: `### ` does not match `## `, and nothing downstream noticed.
            if s:
                orphaned.append(s)
        elif s.lower().startswith("assessment:"):
            cur["assessment"] = s.split(":", 1)[1].strip()
            last_field = "assessment"
        elif s.lower().startswith("mitigant:"):
            cur["mitigant"] = s.split(":", 1)[1].strip()
            last_field = "mitigant"
        elif s.startswith(SCAFFOLD_PREFIX.strip()):
            cur["scaffold"].append(s[len(SCAFFOLD_PREFIX.strip()) :].strip())
        elif s.startswith(ANALYST_PREFIX.strip()):
            cur["analyst"].append(s[len(ANALYST_PREFIX.strip()) :].strip())
        elif s:
            # A continuation line belongs to whichever field it follows. It used to
            # be appended ONLY when a mitigant already existed, so wrapped prose
            # under Assessment — and anything between the heading and the first
            # labelled field — was discarded with no trace.
            if last_field:
                cur[last_field] += " " + s
            else:
                cur["unparsed"].append(s)
    if cur:
        blocks.append(cur)
    if orphaned and blocks:
        blocks[0]["unparsed"] = orphaned + blocks[0]["unparsed"]
    elif orphaned:
        # Nothing parsed as a risk at all; hand the caller everything so it can
        # fall back to prose rather than rendering an empty table.
        blocks.append(
            {
                "name": "",
                "assessment": "",
                "mitigant": "",
                "scaffold": [],
                "analyst": [],
                "unparsed": orphaned,
            }
        )
    return blocks


def risks_parse_is_usable(blocks: list[dict]) -> bool:
    """True when the parse produced something worth rendering as a table.

    `render_deal` took the table branch whenever ANY `## ` line existed, so a
    Section 7 written as ordinary prose with subheadings rendered as a table of
    em-dashes — the section where the paper's risks live, silently emptied. If no
    block carries an assessment or a mitigant, the agent did not write the risk
    format and the text should render as prose instead.
    """
    return any(b["assessment"] or b["mitigant"] for b in blocks)


def _render_risks_table(
    doc,
    blocks: list[dict],
    spec: PaperSpec,
    wip: _WipRegister | None = None,
    section_label: str | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    _mark_header_row(table.rows[0])
    for i, h in enumerate(("Risk", "Assessment", "Mitigant")):
        table.rows[0].cells[i].text = h
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for b in blocks:
        cells = table.add_row().cells
        cells[0].text = b["name"]
        cells[1].text = b["assessment"] or "—"
        cells[2].text = b["mitigant"] or "—"
    # Scaffold / analyst items render below the table (keeping the table clean) and
    # the analyst actions roll into the WIP register, like every other section.
    # Content the parser could not place must be VISIBLE, never dropped. A risk the
    # agent wrote but the format did not fit is still a risk, and Section 7 is where
    # check #17's disclosures live — silently losing a line here loses a disclosure.
    for b in blocks:
        for line in b.get("unparsed", []):
            _analyst_para(doc, f"[unparsed — confirm placement] {line}", spec)
            if wip is not None:
                wip.add(
                    section_label or "7. Risks & Mitigations",
                    f"Place unparsed risk content: {line[:120]}",
                )

    extras = [b for b in blocks if b["scaffold"] or b["analyst"]]
    if extras:
        h = doc.add_paragraph()
        h.add_run("Outstanding to confirm").bold = True
        for b in extras:
            for line in b["scaffold"]:
                _scaffold_para(doc, f"{b['name']}: {line}", spec)
            for action in b["analyst"]:
                _analyst_para(doc, f"{b['name']} — {action}", spec)
                if wip is not None:
                    wip.add(section_label or "7. Risks & Mitigations", f"{b['name']}: {action}")
            if b["scaffold"] and not b["analyst"] and wip is not None:
                wip.add(section_label or "7. Risks & Mitigations", f"Complete: {b['name']}")


def _render_sources(doc, model: DealModel, block, prov: _Provenance, spec: PaperSpec) -> None:
    seen: list[str] = []
    for path in block.cited_field_paths:
        target = resolve_path(model, path)
        m = prov.marker(target if isinstance(target, Fact) else None).strip()
        if m and m not in seen:
            seen.append(m)
    if seen:
        src = doc.add_paragraph()
        run = src.add_run("Sources: " + " ".join(seen))
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(*spec.muted_rgb)


def _scaffold_para(doc, text: str, spec: PaperSpec) -> None:
    """A best-practice template line to be completed — distinct draft styling so it
    reads unmistakably as scaffold, not finished prose."""
    bullet = text.lstrip().startswith(("- ", "• "))
    body = text.lstrip()[2:].strip() if bullet else text
    p = doc.add_paragraph(style="List Bullet") if bullet else doc.add_paragraph()
    run = p.add_run(body)
    run.italic = True
    run.font.color.rgb = RGBColor(*spec.draft_rgb)


def _analyst_para(doc, text: str, spec: PaperSpec) -> None:
    """An explicit human action, tagged distinctly and rolled into the WIP register."""
    p = doc.add_paragraph()
    tag = p.add_run("Analyst to confirm: ")
    tag.bold = True
    tag.font.color.rgb = RGBColor(*spec.draft_rgb)
    body = p.add_run(text)
    body.italic = True
    body.font.color.rgb = RGBColor(*spec.draft_rgb)


def _md_cells(line: str) -> list[str]:
    """Cells of one Markdown pipe-table row (drop the outer empty edges)."""
    parts = [c.strip() for c in line.strip().strip("|").split("|")]
    return parts


def _is_md_separator(line: str) -> bool:
    """The '| --- | --- |' alignment row beneath a table header."""
    cells = _md_cells(line)
    return bool(cells) and all(set(c) <= set("-: ") and "-" in c for c in cells)


def _render_md_table(doc, rows: list[str], spec: PaperSpec) -> None:
    """Render a contiguous Markdown pipe-table (header + separator + body) as a
    native Word table (D2: plain, flat, Canva-importable)."""
    parsed = [_md_cells(r) for r in rows if not _is_md_separator(r)]
    parsed = [r for r in parsed if any(c for c in r)]
    if not parsed:
        return
    ncols = max(len(r) for r in parsed)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Light Grid Accent 1"
    _mark_header_row(table.rows[0])
    for i in range(ncols):
        table.rows[0].cells[i].text = parsed[0][i] if i < len(parsed[0]) else ""
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for r in parsed[1:]:
        cells = table.add_row().cells
        for i in range(ncols):
            cells[i].text = r[i] if i < len(r) else ""


def _render_prose(
    doc,
    model: DealModel,
    section: int,
    block,
    prov: _Provenance,
    show: bool,
    spec: PaperSpec,
    wip: _WipRegister | None = None,
    section_label: str | None = None,
) -> None:
    sub = subsub = 0
    cur_heading = section_label or f"Section {section}"
    scaffold_no_action: set[str] = set()  # headings with scaffold but no analyst line
    lines = block.text.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        # A contiguous run of '|' lines is a Markdown pipe-table -> native table.
        if s.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i])
                i += 1
            _render_md_table(doc, tbl, spec)
            continue
        i += 1
        if s.startswith("### "):
            subsub += 1
            cur_heading = s[4:].strip()
            doc.add_heading(f"{section}.{sub}.{subsub}  {cur_heading}", level=3)
        elif s.startswith("## "):
            sub += 1
            subsub = 0
            cur_heading = s[3:].strip()
            doc.add_heading(f"{section}.{sub}  {cur_heading}", level=2)
        elif s.startswith(SCAFFOLD_PREFIX.strip()):
            _scaffold_para(doc, s[len(SCAFFOLD_PREFIX.strip()) :].strip(), spec)
            scaffold_no_action.add(cur_heading)
        elif s.startswith(ANALYST_PREFIX.strip()):
            action = s[len(ANALYST_PREFIX.strip()) :].strip()
            _analyst_para(doc, action, spec)
            scaffold_no_action.discard(cur_heading)
            if wip is not None:
                wip.add(section_label or cur_heading, action)
        else:
            doc.add_paragraph(s)
    # A scaffolded subsection with no explicit analyst line still owes completion.
    if wip is not None:
        for heading in scaffold_no_action:
            wip.add(section_label or f"Section {section}", f"Complete: {heading}")
    if show and block.cited_field_paths:
        _render_sources(doc, model, block, prov, spec)
    if section == 9:
        # D8, visibility half. The structural gate refuses APPROVED, but that fires
        # long after the draft and says nothing to whoever reads the .docx. Surface
        # it here so the analyst sees the offending sentence in the register rather
        # than discovering it when approval is refused — and, more importantly, so
        # a reader of the draft is told that the prose above the recommendation
        # line contains one.
        if wip is not None:
            for sentence in machine_recommendations(block.text):
                wip.add(
                    section_label or "9. Conclusion",
                    f"REMOVE — the drafted body states a credit recommendation, "
                    f"which is human-authored (D8): {sentence!r}",
                )
        rec = doc.add_paragraph()
        if block.human_recommendation and block.human_recommendation.strip():
            rec.add_run("Recommendation. ").bold = True
            rec.add_run(block.human_recommendation.strip())
        else:
            run = rec.add_run(
                "[Recommendation — to be authored by the analyst; "
                "human-authored and required before approval (D8).]"
            )
            run.italic = True
            run.font.color.rgb = RGBColor(*spec.draft_rgb)
            if wip is not None:
                wip.add(
                    section_label or "9. Conclusion", "Author the credit recommendation (human-authored, D8)"
                )


# ---------------------------------------------------------------------------
# Document
# ---------------------------------------------------------------------------


def _render_approval(
    doc,
    spec: PaperSpec,
    committee: list[tuple[str, str]] | None,
    *,
    approval: dict | None = None,
) -> None:
    """Final signing page (the signed paper has one; D8 — committee sign-off)."""
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Credit Approval", level=1)
    if approval:
        actor_name = str(approval.get("actor_name") or "Former member")
        approved_at = approval.get("approved_at")
        approved_at_text = (
            approved_at.isoformat()
            if hasattr(approved_at, "isoformat")
            else str(approved_at or "unknown date")
        )
        doc.add_paragraph(f"Approved by {actor_name} · {approved_at_text}")
        if approval.get("id"):
            doc.add_paragraph(f"Approval reference: {approval['id']}")
        if approval.get("digest"):
            doc.add_paragraph(f"Approval digest: {approval['digest']}")
        if approval.get("ledger_revision") is not None:
            doc.add_paragraph(f"Ledger revision: {approval['ledger_revision']}")
    p = doc.add_paragraph()
    p.add_run(
        "We, the undersigned members of the Credit Committee, confirm that we have "
        "reviewed this Internal Credit Memorandum and the supporting documentation; "
        "that we understand the material risks of the proposed transaction; that we "
        "accept those risks as consistent with the firm's credit policy and risk "
        "tolerance; and that we approve the facility on the terms set out herein."
    )
    note = doc.add_paragraph()
    nrun = note.add_run(
        "This page constitutes formal approval and is retained as part of the Credit "
        "Committee's records. The Section 9 recommendation must be human-authored "
        "before approval (D8)."
    )
    nrun.italic = True
    nrun.font.size = Pt(9)
    nrun.font.color.rgb = RGBColor(*spec.muted_rgb)

    rows = committee if committee else ([(actor_name, "Approver")] if approval else [("", "")] * 4)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    _mark_header_row(table.rows[0])
    for i, h in enumerate(("Name", "Role", "Signature", "Date")):
        table.rows[0].cells[i].text = h
        table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for name, role in rows:
        cells = table.add_row().cells
        if name:
            cells[0].text = name
            cells[1].text = role
        else:
            _placeholder(cells[0], "committee member", spec)
        if approval:
            cells[2].text = "Digitally approved"
            cells[3].text = approved_at_text
        else:
            cells[2].text = ""  # signature
            cells[3].text = ""  # date
    if not committee and not approval:
        cm = doc.add_paragraph()
        crun = cm.add_run("[Committee membership to be populated per client configuration.]")
        crun.italic = True
        crun.font.size = Pt(9)
        crun.font.color.rgb = RGBColor(*spec.draft_rgb)


def _render_wip(doc, wip: _WipRegister, spec: PaperSpec) -> None:
    """The 'Items to be Confirmed' register: every outstanding item and analyst
    action, grouped by section, so the paper's level of completion is explicit.
    Rendered as an annexure at the very end (Jason 2026-06-16) so it does not
    interrupt the flow from the Conclusion to the sign-off."""
    if not wip.items:
        return
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _heading(doc, "Annexure — Items to be Confirmed", 1, spec, draft=True)
    note = doc.add_paragraph()
    nrun = note.add_run(
        "Outstanding items and analyst actions identified in this draft. The paper "
        "is not final until each is resolved or deliberately accepted on file."
    )
    nrun.italic = True
    nrun.font.size = Pt(9)
    nrun.font.color.rgb = RGBColor(*spec.muted_rgb)
    groups: dict[str, list[str]] = {}
    for label, item in wip.items:
        groups.setdefault(label, []).append(item)
    for label, items in groups.items():
        h = doc.add_paragraph()
        h.add_run(label).bold = True
        for item in items:
            b = doc.add_paragraph(style="List Bullet")
            run = b.add_run(item)
            run.font.color.rgb = RGBColor(*spec.draft_rgb)


def _header(
    doc,
    model: DealModel,
    paper_title: str,
    spec: PaperSpec,
    *,
    status: str = "DRAFT",
    version_reference: str | None = None,
) -> None:
    """Deal name + paper reference on every page (comment 1, 2026-06-16)."""
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.text = ""
    name = str(model.deal.name.value or model.deal_id)
    ref = f"{paper_title} · {status}"
    if version_reference:
        ref = f"{ref} · {version_reference}"
    # Avoid repeating the name when the deal id is just the name (Phase-0 default).
    if model.deal_id and model.deal_id != name:
        ref = f"{model.deal_id} · {ref}"
    if spec.logo_bytes:
        # Inline picture only (D2): no floating shapes in the header.
        _add_logo(p.add_run(), spec.logo_bytes, height_mm=8)
        p.add_run("    ")
    run = p.add_run(f"{name}    {ref}")
    run.font.size = Pt(spec.caption_size)
    run.font.color.rgb = RGBColor(*spec.muted_rgb)


def _add_logo(run, logo_bytes: bytes, *, height_mm: float) -> bool:
    """Add the firm logo as an inline picture; a logo Word cannot read is skipped."""
    try:
        picture = run.add_picture(io.BytesIO(logo_bytes), height=Mm(height_mm))
        picture._inline.docPr.set("descr", "Firm logo")
        return True
    except Exception:  # noqa: BLE001 — a bad logo never fails a render
        return False


def _render_cover(
    doc,
    model: DealModel,
    paper_title: str,
    spec: PaperSpec,
    paper_date: date,
    *,
    status: str = "DRAFT",
    version_reference: str | None = None,
    approval_id: str | None = None,
    include_firm_name_with_logo: bool = False,
) -> None:
    if spec.logo_bytes:
        _add_logo(doc.add_paragraph().add_run(), spec.logo_bytes, height_mm=18)
    if spec.firm_name and (not spec.logo_bytes or include_firm_name_with_logo):
        firm = doc.add_paragraph().add_run(spec.firm_name)
        firm.bold = True
        firm.font.name = spec.heading_font
        firm.font.color.rgb = RGBColor(*spec.heading_rgb)
    doc.add_heading(str(model.deal.name.value or model.deal_id), level=0)
    sub = doc.add_paragraph()
    sub.add_run(paper_title).bold = True
    dtype = model.deal.type.value.value.replace("_", " ").title() if model.deal.type.value else "—"
    details = [
        ("Deal type", dtype),
        ("Subject property", str(model.valuation.property.address.value or "—")),
        ("Status", status),
        ("Version", version_reference or f"model v{model.model_version}"),
    ]
    if status != "DRAFT":
        details.append(("Approval reference", approval_id or "—"))
    details.append(("Date", _fmt_date(paper_date)))
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, val in details:
        c = t.add_row().cells
        c[0].text = k
        c[1].text = val
    if status == "DRAFT":
        note = doc.add_paragraph()
        run = note.add_run(
            "DRAFT — generated; pending analyst review. Sections marked [DRAFT] and "
            "placeholders in red are to be completed as information is confirmed."
        )
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(*spec.draft_rgb)
    else:
        note = doc.add_paragraph()
        run = note.add_run(f"{status} — frozen approved version. Approval reference: {approval_id or '—'}.")
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(*spec.muted_rgb)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def render_deal(
    model: DealModel,
    out_path: str | Path,
    provenance_mode: str = "appendix",
    paper_title: str = "Internal Credit Memorandum",
    spec: PaperSpec = DEFAULT_SPEC,
    paper_date: date | None = None,
    committee: list[tuple[str, str]] | None = None,
    house_style: dict | None = None,
) -> Path:
    doc = Document()
    apply_paper_spec(doc, spec)
    prov = _Provenance()
    wip = _WipRegister()
    section_titles = _section_titles(house_style)
    show = provenance_mode == "appendix"

    if spec.running_header:
        _header(doc, model, paper_title, spec)
    _render_cover(doc, model, paper_title, spec, paper_date or date.today())

    for section in range(1, 10):
        title = section_titles.get(section, f"Section {section}")
        label = f"{section}. {title}"
        block = model.prose.get(section)
        is_draft = block is None or block.status != SectionStatus.APPROVED
        _heading(doc, f"{section}.  {title}", 1, spec, draft=is_draft)
        # Say who wrote it (DECISIONS.md 2026-08-21). A reader must be able to tell
        # engine prose from a person's, and cannot infer it from anything else on
        # the page — an APPROVED section is still AI-authored unless a human
        # rewrote it. Sections 1 and 3 are tables built from anchored facts, so
        # they are marked deterministic rather than generated.
        if block is not None:
            _authorship_note(doc, block.authorship, spec)
        if section == 1:
            _render_snapshot(doc, model, prov, show, spec, wip, label)
        elif section == 3:
            _render_funding(doc, model, prov, show, spec)
        elif block is not None and block.text.strip():
            if section == 7 and (blocks := _parse_risks(block.text)) and risks_parse_is_usable(blocks):
                _render_risks_table(doc, blocks, spec, wip, label)
                if show and block.cited_field_paths:
                    _render_sources(doc, model, block, prov, spec)
            else:
                _render_prose(doc, model, section, block, prov, show, spec, wip, label)
        else:
            # Section not yet generated — boilerplate placeholder.
            p = doc.add_paragraph()
            run = p.add_run(f"[{title} — to be generated and populated.]")
            run.italic = True
            run.font.color.rgb = RGBColor(*spec.draft_rgb)
            wip.add(label, "Generate and populate this section")
            if section == 9:
                rec = doc.add_paragraph()
                rrun = rec.add_run(
                    "[Recommendation — to be authored by the analyst; "
                    "human-authored and required before approval (D8).]"
                )
                rrun.italic = True
                rrun.font.color.rgb = RGBColor(*spec.draft_rgb)

    if show and prov.entries:
        _heading(doc, "Provenance", 1, spec)
        ap = doc.add_table(rows=1, cols=2)
        ap.style = "Light Grid Accent 1"
        _mark_header_row(ap.rows[0])
        ap.rows[0].cells[0].text = "Ref"
        ap.rows[0].cells[1].text = "Source"
        for c in ap.rows[0].cells:
            for run in c.paragraphs[0].runs:
                run.bold = True
        for n, p in prov.entries:
            row = ap.add_row().cells
            row[0].text = f"[{n}]"
            page = f", p.{p.page}" if p.page else ""
            if p.method == ExtractionMethod.DERIVED:
                row[1].text = f"Derived: {p.location}"
            elif p.method == ExtractionMethod.HUMAN_ENTRY:
                row[1].text = f"{p.document_id}{page} (analyst-entered)"
            else:
                row[1].text = f"{p.document_id}{page}"

    _render_approval(doc, spec, committee)

    # Outstanding-items register last of all — an annexure that does not interrupt
    # the flow from Conclusion to sign-off (Jason 2026-06-16).
    _render_wip(doc, wip, spec)

    if spec.table_size:
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for r in paragraph.runs:
                            r.font.size = Pt(spec.table_size)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
