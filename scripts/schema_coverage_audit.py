"""Schema coverage audit: what do signed credit papers cite that the deal model cannot hold?

A credit paper is built from facts. Ingestion can only return paths the field
reference offers, so a figure the paper always reports but the schema has no
field for is missed by every model on every deal: measured 2026-10-08, most of
the signed paper's figures that ingestion missed sat in documents it had read,
under labels no concept knows.

For every figure line in every signed paper (amount, percentage, date, count),
this takes the LABEL written before the figure, normalises it, and keeps only
labels that recur across at least --min-papers DIFFERENT papers, and drops any
label naming a party (company suffixes, trusts, signing envelopes): a deal's
names and addresses appear in one paper, but a lender's own entity appears in
every paper it writes, so recurrence alone is not enough. Each recurring
label is then matched against the concept registry's aliases and the deal
model's field names.

PRINTS only recurring labels (shared vocabulary of credit papers) and counts.
Reads papers from fixtures/ (tenant data); writes nothing.

Usage:
    python scripts/schema_coverage_audit.py [--fixtures DIR] [--min-papers 3]
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

# A label naming a party is never a field and may be a name: every paper a lender
# writes repeats its own entity, so recurrence alone does not make a label generic.
_PARTY = re.compile(
    r"\b(pty|ltd|limited|acn|abn|trust|trustee|docusign|envelope|capital|financial|holdings)\b"
)
_STOP = {"the", "of", "and", "a", "an", "to", "in", "for", "on", "at", "by", "is", "as", "per", "with", "be"}


def paper_texts(fixtures: Path) -> dict[str, str]:
    """Every signed paper's text, keyed by an anonymous label."""
    from creditpaper.ingest.reader import PdfiumReader

    sources = sorted((fixtures / "prior_papers").rglob("*.pdf")) + sorted(
        (fixtures / "prior_papers").rglob("*.docx")
    )
    sources += sorted(p for p in (fixtures / "blind").glob("*/*.pdf"))
    sources += sorted(p for p in (fixtures / "build_deal").glob("*.pdf"))
    out: dict[str, str] = {}
    for i, path in enumerate(sources, start=1):
        try:
            if path.suffix.lower() == ".docx":
                from docx import Document

                doc = Document(str(path))
                parts = [p.text for p in doc.paragraphs]
                for table in doc.tables:
                    for row in table.rows:
                        parts.append("  ".join(c.text for c in row.cells))
                text = "\n".join(parts)
            else:
                with PdfiumReader(path) as reader:
                    text = "\n".join(reader.page(n).text for n in range(1, reader.page_count() + 1))
        except Exception:
            continue
        if text.strip():
            out[f"paper{i}"] = text
    return out


def labels(text: str) -> set[str]:
    """The normalised label before the first figure on each line that has one."""
    from creditpaper.ingest.values import figures

    found: set[str] = set()
    for line in text.splitlines():
        figs = [
            f
            for f in figures(line)
            if len(re.sub(r"\D", "", f.value)) >= 3 or f.kind.value in ("percent", "date")
        ]
        if not figs:
            continue
        before = line[: figs[0].start]
        words = [w for w in re.findall(r"[a-z]+", before.casefold()) if w not in _STOP and len(w) > 1]
        label = " ".join(words[-5:])
        if 1 <= len(words) <= 8 and not _PARTY.search(label):
            found.add(label)
    return found


def vocabulary() -> list[tuple[str, str]]:
    """(phrase, path) for every concept alias and every model leaf name."""
    from creditpaper.ingest.harvest_prompt import _leaf_paths, normalise_key_path
    from creditpaper.model.concepts import REGISTRY
    from creditpaper.paths import field_reference

    out = [(a, c.path) for c in REGISTRY.values() for a in c.aliases]
    for path in _leaf_paths(field_reference().read_text(encoding="utf-8")):
        leaf = re.sub(r"\[[^\]]*\]", "", normalise_key_path(path)).split(".")[-1]
        out.append((leaf.replace("_", " "), re.sub(r"\[[^\]]*\]", "[*]", path)))
    return out


def best_match(label: str, vocab: list[tuple[str, str]]) -> tuple[str | None, int]:
    from rapidfuzz import fuzz

    best: tuple[str | None, int] = (None, 0)
    for phrase, path in vocab:
        score = int(fuzz.token_set_ratio(label, phrase))
        if score > best[1]:
            best = (path, score)
    return best


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--fixtures", type=Path, default=ROOT.parent / "creditpaper" / "fixtures")
    ap.add_argument("--min-papers", type=int, default=3)
    ap.add_argument("--match", type=int, default=90, help="token-set similarity counted as covered")
    args = ap.parse_args(argv)

    papers = paper_texts(args.fixtures)
    in_papers: dict[str, set[str]] = defaultdict(set)
    for name, text in papers.items():
        for label in labels(text):
            in_papers[label].add(name)
    recurring = {lab: len(names) for lab, names in in_papers.items() if len(names) >= args.min_papers}
    vocab = vocabulary()
    covered, uncovered = [], []
    for label, n in sorted(recurring.items(), key=lambda kv: -kv[1]):
        path, score = best_match(label, vocab)
        (covered if score >= args.match else uncovered).append((n, label, path, score))
    print(f"{len(papers)} papers; {len(recurring)} figure labels recur in >= {args.min_papers} papers")
    print(f"\nCOVERED ({len(covered)}): label -> field")
    for n, label, path, _score in covered:
        print(f"  {n:>2}  {label:<40} {path}")
    print(f"\nNOT COVERED ({len(uncovered)}): label  (closest field, similarity)")
    for n, label, path, score in uncovered:
        print(f"  {n:>2}  {label:<40} ({path}, {score})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
